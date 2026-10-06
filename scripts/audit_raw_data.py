"""Read-only full Excel audit; preserve candidate rows with Excel row provenance.

Run from any directory. Outputs are local analysis intermediates, never raw edits.
Requires pandas, numpy and openpyxl. No pricing or training is performed.
"""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NUMERIC = ['Bid', 'Ask', 'UnderlyingBid', 'UnderlyingAsk', 'Last',
           'UnderlyingLast', 'ImpliedVolatility', 'Delta', 'Gamma', 'Vega',
           'Theta', 'Volume', 'OpenInterest',
           'ContractSize', 'Strike']


def sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def audit(path, output):
    output.mkdir(parents=True, exist_ok=True)
    before = sha256(path)
    book = load_workbook(path, read_only=True, data_only=True)
    if len(book.sheetnames) != 1:
        raise ValueError('Select and document a sheet before auditing multiple sheets.')
    sheet = book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    columns = list(next(rows))
    profile = {c: Counter() for c in NUMERIC}
    minima, maxima = {}, {}
    categories = {c: Counter() for c in ['SecurityID', 'Currency', 'CallPut',
                  'ExerciseStyle', 'OptionStyle', 'Issuer', 'ContractSize', 'CalculationPrice']}
    years, counters, negative_values = Counter(), Counter(), {c: Counter() for c in NUMERIC}
    date_min = date_max = None
    seen = set()
    candidates = []

    def process(batch, first_row):
        nonlocal date_min, date_max
        df = pd.DataFrame(batch, columns=columns)
        df.insert(0, 'SourceRow', np.arange(first_row, first_row + len(df)))
        empty = df[columns].isna().all(axis=1)
        counters['empty_excel_rows_ignored'] += int(empty.sum())
        df = df.loc[~empty].copy()
        if df.empty:
            return
        counters['records'] += len(df)
        for c in NUMERIC:
            s = pd.to_numeric(df[c], errors='coerce')
            finite = np.isfinite(s)
            profile[c].update({'missing_or_nonnumeric': int(s.isna().sum()),
                               'nonfinite': int((~finite & s.notna()).sum()),
                               'negative': int((s < 0).sum()), 'zero': int((s == 0).sum()),
                               'sentinel_minus_99_99': int(np.isclose(s, -99.99, atol=1e-4, rtol=0).sum())})
            if finite.any():
                minima[c] = min(minima.get(c, float('inf')), float(s[finite].min()))
                maxima[c] = max(maxima.get(c, -float('inf')), float(s[finite].max()))
            negative_values[c].update(s[s < 0].round(5).value_counts().to_dict())
        for c in categories:
            categories[c].update(df[c].fillna('<missing>').astype(str).value_counts().to_dict())
        dates = pd.to_datetime(df.Date, errors='coerce')
        expiry = pd.to_datetime(df.Expiration, errors='coerce')
        years.update(dates.dt.year.dropna().astype(int).value_counts().to_dict())
        counters['missing_date'] += int(dates.isna().sum())
        counters['missing_expiration'] += int(expiry.isna().sum())
        counters['expired_or_same_day'] += int((expiry <= dates).sum())
        if dates.notna().any():
            date_min = min(date_min or dates.min(), dates.min())
            date_max = max(date_max or dates.max(), dates.max())
        for key in zip(df.SecurityID, df.OptionID, dates):
            if key in seen:
                counters['duplicate_security_option_date_excess'] += 1
            seen.add(key)
        bid, ask = pd.to_numeric(df.Bid), pd.to_numeric(df.Ask)
        mid = (bid + ask) / 2
        valid_both = (bid > 0) & (ask > 0)
        counters['both_option_quotes_positive'] += int(valid_both.sum())
        counters['positive_mid'] += int((mid > 0).sum())
        counters['mid_above_0_5'] += int((mid > .5).sum())
        counters['positive_mid_with_nonpositive_leg'] += int(((mid > 0) & ~valid_both).sum())
        counters['crossed_option_quotes_positive'] += int((valid_both & (ask < bid)).sum())
        counters['locked_option_quotes_positive'] += int((valid_both & (ask == bid)).sum())
        candidates.append(df.loc[(mid > 0) | valid_both].copy())
        print(f'Read {counters["records"]:,} source rows', flush=True)

    batch, start = [], 2
    for row in rows:
        batch.append(row)
        if len(batch) == 50000:
            process(batch, start)
            start += len(batch)
            batch = []
    if batch:
        process(batch, start)
    book.close()
    after = sha256(path)
    if after != before:
        raise RuntimeError('Input changed during audit.')
    candidate_df = pd.concat(candidates, ignore_index=True)
    candidate_df.to_csv(output / 'raw_candidates.csv', index=False)
    result = {'source': str(path.resolve()), 'sha256': before,
              'source_unchanged': before == after, 'sheet': sheet.title,
              'columns': columns, 'counters': dict(counters),
              'date_min': str(date_min.date()), 'date_max': str(date_max.date()),
              'year_counts': dict(sorted(years.items())),
              'numeric_profile': {c: {**profile[c], 'min': minima.get(c),
                      'max': maxima.get(c), 'top_negative_values': negative_values[c].most_common(5)} for c in NUMERIC},
              'categories': {c: dict(v) for c, v in categories.items()},
              'candidate_rows': len(candidate_df)}
    (output / 'raw_profile.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps({**result['counters'], 'candidates': len(candidate_df),
                      'source_unchanged': True}, indent=2), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=ROOT / 'data/raw/usal9td17mj9mspt.xlsx')
    parser.add_argument('--output', type=Path, default=ROOT / 'local_metadata/data_audit')
    args = parser.parse_args()
    audit(args.input, args.output)
